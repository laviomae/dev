#!/usr/bin/env python3
"""
SEAO - Convertisseur JSON vers Excel (2024-2026)

Télécharge tous les fichiers JSON du SEAO pour 2024-2026 depuis donneesquebec.ca
et les convertit en un fichier Excel structuré.

Installation:
    pip install requests pandas openpyxl tqdm

Usage:
    python seao_to_excel.py
    python seao_to_excel.py --output mon_fichier.xlsx
    python seao_to_excel.py --cache-only   # si fichiers déjà téléchargés
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
import pandas as pd
from tqdm import tqdm


DATASET_ID = "systeme-electronique-dappel-doffres-seao"
CKAN_API = f"https://www.donneesquebec.ca/recherche/api/3/action/package_show?id={DATASET_ID}"
CACHE_DIR = Path("seao_cache")
OUTPUT_FILE = "seao_2024_2026.xlsx"

HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; SEAO-Analyzer/1.0)"}


# ─────────────────────────────────────────────
# 1. Récupérer et filtrer la liste des fichiers
# ─────────────────────────────────────────────
def fetch_resource_list():
    print("Récupération de la liste des fichiers depuis donneesquebec.ca...")
    resp = requests.get(CKAN_API, headers=HEADERS, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    if not data.get("success"):
        raise RuntimeError(f"Erreur API CKAN: {data}")
    return data["result"]["resources"]


def filter_json_2024_2026(resources):
    """
    Garde les fichiers JSON dont la période couvre 2024-2026.
    Si deux fichiers couvrent la même période, garde le plus récent.
    """
    date_pattern = re.compile(r"(\d{8})_(\d{8})\.json", re.IGNORECASE)

    candidates = []
    for r in resources:
        if r.get("format", "").upper() != "JSON":
            continue
        url = r.get("url", "")
        name = r.get("name", "") or os.path.basename(url)
        m = date_pattern.search(name) or date_pattern.search(url)
        if not m:
            continue
        start_str, end_str = m.group(1), m.group(2)
        # Exclure si entièrement hors de 2024-2026
        if int(end_str[:4]) < 2024 or int(start_str[:4]) > 2026:
            continue
        r["_start"] = start_str
        r["_end"] = end_str
        r["_name"] = f"{start_str}_{end_str}.json"
        candidates.append(r)

    # Dédupliquer par période : garder le plus récent
    def mod_date(r):
        for field in ("last_modified", "metadata_modified", "created"):
            v = r.get(field)
            if v:
                try:
                    return datetime.fromisoformat(v.replace("Z", "+00:00"))
                except Exception:
                    pass
        return datetime.min

    period_map = {}
    for r in candidates:
        key = (r["_start"], r["_end"])
        if key not in period_map or mod_date(r) > mod_date(period_map[key]):
            period_map[key] = r

    result = sorted(period_map.values(), key=lambda x: x["_start"])
    print(f"  → {len(result)} fichiers JSON sélectionnés pour 2024-2026")
    return result


# ─────────────────────────────────────────────
# 2. Téléchargement avec cache
# ─────────────────────────────────────────────
def download_files(resources):
    CACHE_DIR.mkdir(exist_ok=True)
    local_paths = []
    for r in tqdm(resources, desc="Téléchargement"):
        # Utiliser le nom original de l'URL
        filename = os.path.basename(r["url"])
        cache_path = CACHE_DIR / filename
        if cache_path.exists():
            local_paths.append(cache_path)
            continue
        try:
            resp = requests.get(r["url"], headers=HEADERS, timeout=180, stream=True)
            resp.raise_for_status()
            with open(cache_path, "wb") as f:
                for chunk in resp.iter_content(65536):
                    f.write(chunk)
            local_paths.append(cache_path)
            time.sleep(0.3)
        except Exception as e:
            print(f"\n  ⚠ Erreur téléchargement {filename}: {e}")
    return local_paths


# ─────────────────────────────────────────────
# 3. Parsing OCDS
# ─────────────────────────────────────────────
def safe(obj, *keys, default=""):
    for k in keys:
        if not isinstance(obj, dict):
            return default
        obj = obj.get(k)
        if obj is None:
            return default
    return obj if obj is not None else default


def party_name(ref, parties):
    """Résout l'ID d'une partie vers son nom."""
    if not ref:
        return ""
    pid = ref.get("id", "") if isinstance(ref, dict) else str(ref)
    for p in parties:
        if p.get("id") == pid:
            return p.get("name", pid)
    return pid


def parse_release(release):
    """
    Extrait TOUTES les lignes d'un release OCDS SEAO.
    Retourne une liste de dicts (une ligne = une soumission ou un contrat).

    Structure OCDS utilisée par le SEAO:
      release.ocid              → identifiant unique du processus
      release.date              → date de publication
      release.buyer             → organisme acheteur
      release.tender            → avis d'appel d'offres
      release.bids.details[]    → soumissions reçues (avec montants par soumissionnaire)
      release.awards[]          → attributions de contrat
      release.parties[]         → répertoire des organisations
    """
    parties = release.get("parties", [])
    ocid = release.get("ocid", "")
    release_date = safe(release, "date")[:10] if safe(release, "date") else ""
    tags = ", ".join(release.get("tag", []))

    # Acheteur
    buyer = party_name(release.get("buyer"), parties)

    # Appel d'offres
    tender = release.get("tender", {})
    tender_id       = safe(tender, "id")
    tender_title    = safe(tender, "title")
    tender_desc     = safe(tender, "description")
    tender_method   = safe(tender, "procurementMethodDetails") or safe(tender, "procurementMethod")
    tender_status   = safe(tender, "status")
    tender_value    = safe(tender, "value", "amount")
    tender_currency = safe(tender, "value", "currency") or "CAD"
    tender_start    = (safe(tender, "tenderPeriod", "startDate") or "")[:10]
    tender_end      = (safe(tender, "tenderPeriod", "endDate") or "")[:10]
    tender_nb_lots  = len(tender.get("lots", []))

    # Classification / catégorie
    items = tender.get("items", [])
    categories = "; ".join(
        safe(it, "classification", "description")
        for it in items if safe(it, "classification", "description")
    )
    cpv_codes = "; ".join(
        safe(it, "classification", "id")
        for it in items if safe(it, "classification", "id")
    )

    # Gagnants (awards)
    award_winners = []
    award_amount = ""
    for award in release.get("awards", []):
        if award.get("status") in ("active", ""):
            for s in award.get("suppliers", []):
                award_winners.append(party_name(s, parties))
            award_amount = safe(award, "value", "amount") or award_amount
    gagnant_str = "; ".join(award_winners)

    rows = []

    # Une ligne par soumission (bid)
    bids = safe(release, "bids", "details") if isinstance(release.get("bids"), dict) else []
    if not isinstance(bids, list):
        bids = []

    for bid in bids:
        bid_id      = safe(bid, "id")
        bid_status  = safe(bid, "status")
        bid_rank    = bid.get("rank", "")
        bid_amount  = safe(bid, "value", "amount")
        bid_curr    = safe(bid, "value", "currency") or "CAD"
        bid_tenderers = bid.get("tenderers", [])
        # Critères d'évaluation / notes
        eval_criteria = "; ".join(
            f"{c.get('title','')}: {c.get('value','')}"
            for c in bid.get("evaluationCriteria", []) if c.get("title")
        )

        for tenderer_ref in bid_tenderers:
            rows.append({
                # Identification
                "ocid":                     ocid,
                "Date publication":         release_date,
                "Tags (type d'avis)":       tags,
                # Acheteur
                "Organisme acheteur":       buyer,
                # Appel d'offres
                "Numéro appel d'offres":    tender_id,
                "Titre":                    tender_title,
                "Description":              tender_desc,
                "Méthode d'appel d'offres": tender_method,
                "Statut appel d'offres":    tender_status,
                "Valeur estimée ($)":       tender_value,
                "Devise":                   tender_currency,
                "Date ouverture":           tender_start,
                "Date clôture":             tender_end,
                "Catégorie":                categories,
                "Code CPV":                 cpv_codes,
                "Nb lots":                  tender_nb_lots if tender_nb_lots else "",
                # Soumission
                "ID soumission":            bid_id,
                "Soumissionnaire":          party_name(tenderer_ref, parties),
                "Montant soumis ($)":       bid_amount,
                "Devise soumission":        bid_curr,
                "Statut soumission":        bid_status,
                "Rang":                     bid_rank,
                "Critères évaluation":      eval_criteria,
                # Attribution
                "Gagnant(s)":               gagnant_str,
                "Montant contrat attribué ($)": award_amount,
            })

    # Si aucun bid mais il y a une attribution, créer quand même des lignes
    if not rows and release.get("awards"):
        for award in release.get("awards", []):
            if award.get("status") not in ("active", ""):
                continue
            a_amount = safe(award, "value", "amount")
            for supplier in award.get("suppliers", []):
                rows.append({
                    "ocid":                     ocid,
                    "Date publication":         release_date,
                    "Tags (type d'avis)":       tags,
                    "Organisme acheteur":       buyer,
                    "Numéro appel d'offres":    tender_id,
                    "Titre":                    tender_title,
                    "Description":              tender_desc,
                    "Méthode d'appel d'offres": tender_method,
                    "Statut appel d'offres":    tender_status,
                    "Valeur estimée ($)":       tender_value,
                    "Devise":                   tender_currency,
                    "Date ouverture":           tender_start,
                    "Date clôture":             tender_end,
                    "Catégorie":                categories,
                    "Code CPV":                 cpv_codes,
                    "Nb lots":                  tender_nb_lots if tender_nb_lots else "",
                    "ID soumission":            "",
                    "Soumissionnaire":          "",
                    "Montant soumis ($)":       "",
                    "Devise soumission":        "CAD",
                    "Statut soumission":        "",
                    "Rang":                     "",
                    "Critères évaluation":      "",
                    "Gagnant(s)":               party_name(supplier, parties),
                    "Montant contrat attribué ($)": a_amount,
                })

    # Si toujours vide (avis sans soumissions ni attributions), une ligne minimale
    if not rows:
        rows.append({
            "ocid":                     ocid,
            "Date publication":         release_date,
            "Tags (type d'avis)":       tags,
            "Organisme acheteur":       buyer,
            "Numéro appel d'offres":    tender_id,
            "Titre":                    tender_title,
            "Description":              tender_desc,
            "Méthode d'appel d'offres": tender_method,
            "Statut appel d'offres":    tender_status,
            "Valeur estimée ($)":       tender_value,
            "Devise":                   tender_currency,
            "Date ouverture":           tender_start,
            "Date clôture":             tender_end,
            "Catégorie":                categories,
            "Code CPV":                 cpv_codes,
            "Nb lots":                  tender_nb_lots if tender_nb_lots else "",
            "ID soumission":            "",
            "Soumissionnaire":          "",
            "Montant soumis ($)":       "",
            "Devise soumission":        "CAD",
            "Statut soumission":        "",
            "Rang":                     "",
            "Critères évaluation":      "",
            "Gagnant(s)":               gagnant_str,
            "Montant contrat attribué ($)": award_amount,
        })

    return rows


def parse_json_file(path):
    rows = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        print(f"\n  ⚠ Erreur lecture {path.name}: {e}")
        return rows

    # Supporter les deux formats OCDS : package {"releases":[...]} ou liste directe
    if isinstance(data, dict):
        releases = (data.get("releases")
                    or data.get("records")
                    or ([data] if "ocid" in data else []))
    elif isinstance(data, list):
        releases = data
    else:
        return rows

    for item in releases:
        if not isinstance(item, dict):
            continue
        # Certains formats encapsulent dans "compiledRelease"
        release = item.get("compiledRelease", item)
        rows.extend(parse_release(release))

    return rows


# ─────────────────────────────────────────────
# 4. Écriture Excel
# ─────────────────────────────────────────────
def build_excel(all_rows, output_file):
    if not all_rows:
        print("⚠ Aucune donnée extraite — vérifiez les fichiers JSON.")
        return

    df = pd.DataFrame(all_rows)

    # Convertir les colonnes numériques
    for col in ["Valeur estimée ($)", "Montant soumis ($)", "Montant contrat attribué ($)"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    df = df.sort_values(["Date publication", "ocid", "Rang"], ascending=[False, True, True])

    print(f"Écriture du fichier Excel : {output_file}")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Toutes les données", index=False)

        ws = writer.sheets["Toutes les données"]

        # Largeur automatique des colonnes
        for col_cells in ws.columns:
            max_len = max(
                (len(str(c.value)) for c in col_cells if c.value is not None),
                default=10,
            )
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 80)

        # Figer la première ligne
        ws.freeze_panes = "A2"

        # Mettre en gras les en-têtes
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill("solid", fgColor="1F4E79")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    print(f"\n✓ Fichier créé : {output_file}")
    print(f"  • {len(df):,} lignes")
    print(f"  • {df['ocid'].nunique():,} appels d'offres uniques")
    print(f"  • {df['Soumissionnaire'].nunique():,} soumissionnaires distincts")
    print(f"  • {df['Organisme acheteur'].nunique():,} organismes acheteurs")


# ─────────────────────────────────────────────
# 5. Main
# ─────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="SEAO JSON → Excel (2024-2026)")
    parser.add_argument("--output", default=OUTPUT_FILE,
                        help=f"Nom du fichier Excel (défaut: {OUTPUT_FILE})")
    parser.add_argument("--cache-only", action="store_true",
                        help="Parser uniquement les fichiers déjà dans seao_cache/")
    args = parser.parse_args()

    if args.cache_only:
        CACHE_DIR.mkdir(exist_ok=True)
        local_paths = sorted(CACHE_DIR.glob("*.json"))
        print(f"Mode cache : {len(local_paths)} fichiers JSON locaux")
    else:
        resources = fetch_resource_list()
        filtered = filter_json_2024_2026(resources)
        local_paths = download_files(filtered)

    print(f"\nParsing de {len(local_paths)} fichiers...")
    all_rows = []
    for path in tqdm(local_paths, desc="Parsing"):
        rows = parse_json_file(path)
        all_rows.extend(rows)

    print(f"  → {len(all_rows):,} enregistrements extraits au total")
    build_excel(all_rows, args.output)


if __name__ == "__main__":
    main()
